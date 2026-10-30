#!/usr/bin/env python3
"""Read-only, aggregate-only ADNI T1 / FreeSurfer availability audit.

Checks whether local file trees contain FreeSurfer stats and T1 NIfTI files
that appear linkable to T1 image metadata in an existing local all.xlsx.
Does NOT check image QC or join clinical visits/dates, and does NOT export IDs.
"""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import re
import time

PTID_RE = re.compile(r"(?<![A-Za-z0-9])(\d{3}_S_\d{4,5})(?![A-Za-z0-9])", re.I)
IMAGE_RE = re.compile(r"(?<![A-Za-z0-9])I(\d{4,9})(?!\d)", re.I)
NEEDED = {"aseg.stats", "lh.aparc.stats", "rh.aparc.stats"}


def tokens(s):
    if not s:
        return set(), set()
    s = str(s)
    p = {m.group(1).upper() for m in PTID_RE.finditer(s)}
    i = {"I" + m.group(1) for m in IMAGE_RE.finditer(s)}
    return p, i


def metadata_index(xlsx):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Please install openpyxl in your local environment")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(rows)]
        col = {h: j for j,h in enumerate(headers)}
        keys = ["Subject ID_T1", "Image ID_T1"]
        missing = [h for h in keys if h not in col]
        if missing:
            raise RuntimeError("Missing T1 metadata column(s): " + ", ".join(missing))
        ps, ids = set(), set()
        n = 0
        for row in rows:
            n += 1
            vpid = row[col["Subject ID_T1"]] if len(row)>col["Subject ID_T1"] else None
            vid = row[col["Image ID_T1"]] if len(row)>col["Image ID_T1"] else None
            p, _ = tokens(vpid)
            _, i = tokens(vid)
            # Many ADNI exports put image IDs in bare numeric form or use a leading I.
            if vid is not None and not i:
                sval = str(vid).strip()
                if re.fullmatch(r"\d{4,9}(\.0)?", sval):
                    i.add("I"+sval.split(".")[0])
                elif re.fullmatch(r"[Ii]\d{4,9}", sval):
                    i.add("I"+sval[1:])
            ps.update(p)
            ids.update(i)
        return ps, ids, n
    finally:
        wb.close()


def scan_fs(root, wanted_ptids, deadline, max_entries):
    out = dict(root_present=os.path.isdir(root), scan_complete=True, scanned_entries=0,
               stats_directories=0, complete_stats_triples=0,
               recon_mri_norm_exists=0, recon_mri_aparc_aseg_exists=0,
               recon_lh_white_exists=0, recognized_ptid_stats_directories=0,
               stats_directories_matched_to_xlsx_ptid=0, errors=0)
    matched_ps, observed_ps = set(), set()
    if not out['root_present']:
        return out, matched_ps
    for current, dirs, files in os.walk(root, followlinks=False):
        out['scanned_entries'] += len(dirs) + len(files)
        if out['scanned_entries'] >= max_entries or time.monotonic() > deadline:
            out['scan_complete'] = False
            break
        if os.path.basename(current) == 'stats':
            st = set(files)
            if 'aseg.stats' in st:
                out['stats_directories'] += 1
                if NEEDED <= st:
                    out['complete_stats_triples'] += 1
                base = os.path.dirname(current)
                ps, _ = tokens(os.path.relpath(base, root))
                if ps:
                    out['recognized_ptid_stats_directories'] += 1
                    observed_ps.update(ps)
                    common = ps & wanted_ptids
                    if common:
                        matched_ps.update(common)
                        out['stats_directories_matched_to_xlsx_ptid'] += 1
                if os.path.isfile(os.path.join(base, 'mri', 'norm.mgz')):
                    out['recon_mri_norm_exists'] += 1
                if os.path.isfile(os.path.join(base, 'mri', 'aparc+aseg.mgz')):
                    out['recon_mri_aparc_aseg_exists'] += 1
                if os.path.isfile(os.path.join(base, 'surf', 'lh.white')):
                    out['recon_lh_white_exists'] += 1
            dirs[:] = []  # no useful subtrees beneath a stats directory
    out['unique_recognized_ptids_from_stats_paths'] = len(observed_ps)
    out['unique_ptids_matched_to_xlsx'] = len(matched_ps)
    return out, matched_ps


def scan_t1_nii(root, wanted_ptids, wanted_images, deadline, max_entries):
    out = dict(root_present=os.path.isdir(root), scan_complete=True,
               scanned_entries=0, nifti_files=0, matched_image_id_files=0,
               matched_ptid_files=0, errors=0)
    seen_ps, seen_ids, matched_ps = set(), set(), set()
    if not out['root_present']:
        return out
    for current, dirs, files in os.walk(root, followlinks=False):
        out['scanned_entries'] += len(dirs) + len(files)
        if out['scanned_entries'] >= max_entries or time.monotonic() > deadline:
            out['scan_complete'] = False
            break
        for file in files:
            if not (file.lower().endswith('.nii') or file.lower().endswith('.nii.gz')):
                continue
            out['nifti_files'] += 1
            path = os.path.relpath(os.path.join(current,file),root)
            ps, ids = tokens(path)
            seen_ps.update(ps)
            seen_ids.update(ids)
            if ps & wanted_ptids:
                out['matched_ptid_files'] += 1
                matched_ps.update(ps & wanted_ptids)
            if ids & wanted_images:
                out['matched_image_id_files'] += 1
    out['unique_recognized_ptids_in_path'] = len(seen_ps)
    out['unique_recognized_image_ids_in_path'] = len(seen_ids)
    out['unique_matched_ptids'] = len(matched_ps)
    out['unique_matched_image_ids'] = len(seen_ids & wanted_images)
    return out


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--xlsx', required=True, help='Local all.xlsx; read-only')
    p.add_argument('--root', required=True, help='Mounted ADNI directory')
    p.add_argument('--out', required=True, help='Local aggregate-only JSON')
    p.add_argument('--seconds',type=int,default=900)
    p.add_argument('--max-entries',type=int,default=450000)
    a=p.parse_args()
    started=time.monotonic()
    ps, ids, rows=metadata_index(a.xlsx)
    deadline=time.monotonic()+a.seconds
    root=Path(a.root)
    fs={}
    unions=set()
    for name in ('processed_data','recon_all_new'):
        stats, matched=scan_fs(str(root/name), ps, deadline, a.max_entries)
        fs[name]=stats
        unions |= matched
    nii=scan_t1_nii(str(root/'data'/'T1_nii'),ps,ids,deadline,a.max_entries)
    r={
      'version':'4.0.0',
      'audit_type':'on_disk_t1_and_freesurfer_presence_not_quality_control',
      'xlsx_rows_scanned':rows,
      'xlsx_unique_subject_ids_matching_ADNI_Ptid_pattern':len(ps),
      'xlsx_unique_T1_image_IDs_recognized':len(ids),
      'freesurfer':fs,
      'union_unique_ptids_matched_to_any_freesurfer_stats':len(unions),
      't1_nifti':nii,
      'elapsed_seconds':round(time.monotonic()-started,2),
      'limitations':[
        'No raw IDs, paths, dates or per-participant data are exported.',
        'PTID/image-ID matching depends on canonical ID strings in file/directory names.',
        'A low/zero match may reflect naming conventions rather than absent images.',
        'Files are not opened: no corrupt-file detection, MRI quality assessment or temporal alignment.',
        'FreeSurfer stats are not automatically equivalent to a completed T1 MRI segmentation quality check.',
        'All counts may be partial when a scan is incomplete or capped.'
      ]
    }
    out=Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('Aggregate report written. Completion flags:',
          {k:v['scan_complete'] for k,v in fs.items()},
          'T1_nii:',nii['scan_complete'])
    print('Scanned file/dir entries limited to',a.max_entries,'per scope; max total wall time',a.seconds,'s after reading XLSX.')

if __name__=='__main__':
    main()
