#!/usr/bin/env python3
"""Read-only, aggregate-only audit of a wide ADNI-like XLSX workbook.

Dependencies: openpyxl. Never exports participant-level rows, IDs, dates,
cell samples, or raw data. All worksheet names are omitted from JSON.

Usage:
  python adni_xlsx_audit.py schema --xlsx data/all.xlsx --out results/data_audit/05_schema.json
  python adni_xlsx_audit.py cohort --xlsx data/all.xlsx --sheet 1 --out results/data_audit/06_cohort.json

Limitations: this script inspects clinical table contents only. It does not
match images, assess image quality, validate ADNI diagnostic coding, or infer A+.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import math
import os
import re
import sys
import time
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Missing openpyxl. Try: python -m pip install openpyxl")

VERSION = "1.1.0"

ROLES = {
    "subject": {"RID", "PTID", "SUBJECTID", "SUBJECT", "PARTICIPANTID", "SUBJID"},
    "visit": {"VISCODE2", "VISCODE", "LONIVISCODE", "VISIT", "VISITCODE", "EVENT"},
    "date": {"EXAMDATE", "VISITDATE", "ASSESSDATE", "ASSESSMENTDATE", "SCANDATE", "ACQDATE"},
    "diagnosis": {"DX", "DXBL", "DXCHANGE", "DXCURREN", "DIAGNOSIS", "DIAGNOSISBL", "DIAGNOSISCODE", "DXY"},
    "cognitive": {"MMSE", "MMSCORE", "ADAS13", "ADAS11", "CDRSB", "MOCA", "FAQ", "RAVLTIMMEDIATE", "RAVLTLEARNING", "RAVLTFORGETTING", "RAVLTPCFORGETTING", "TRABSCOR", "TRABSCORBL"},
    "amyloid_or_tau_candidate": {"ABETA", "ABETA40", "ABETA42", "ABETA4240", "ABETA42ABETA40", "AV45", "FBB", "CENTILOID", "CENTILOIDS", "AMYLOID", "AMYLOIDSTATUS", "AMYLOIDPOS", "AMYLOIDPOSITIVE", "TAU", "PTAU", "PTAU181", "PTAU217", "TTAU", "AV1451", "FLORTAUCIPIR"},
    "imaging_link": {"IMAGEUID", "IMAGEUIDLONI", "IMAGEID", "LONISID", "LONIUID", "SCANID", "SERIESUID", "SERIESID"},
    "covariate": {"AGE", "PTGENDER", "GENDER", "SEX", "PTEDUCAT", "EDUCATION", "APOE4", "ORIGPROT", "COLPROT", "SITE", "SITEID"},
}
CANONICAL_TO_ROLE = {f: role for role, fields in ROLES.items() for f in fields}
SCORES = ("MMSE", "ADAS13", "ADAS11", "CDRSB", "MOCA", "FAQ")
SCORE_LIMITS = {"MMSE": (0, 30), "MOCA": (0, 30), "CDRSB": (0, 18), "ADAS13": (0, 85), "ADAS11": (0, 70), "FAQ": (0, 30)}
MISSING_TEXT = {"", "NA", "N/A", "NAN", "NULL", "NONE", "MISSING", "-", "-4", "-9", "-1", "-999"}


def canonical(value):
    if value is None:
        return ""
    s = str(value).strip()
    # pandas or manual merges may suffix duplicate headers with .1, .2, ...
    s = re.sub(r"\.\d+$", "", s)
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def classify(header):
    key = canonical(header)
    role = CANONICAL_TO_ROLE.get(key)
    if role is None:
        if key.endswith("BL") and key[:-2] in CANONICAL_TO_ROLE:
            role = CANONICAL_TO_ROLE[key[:-2]]
        else:
            return None
    # Include only known, whitelisted field identifiers, never arbitrary headers.
    return {"field": key, "role": role}


def scan_schema(workbook, header_row):
    sheets = []
    for sheet_ix, ws in enumerate(workbook.worksheets, 1):
        declared_rows, declared_columns = ws.max_row, ws.max_column
        ws.reset_dimensions()  # Can prevent silently truncated reads from bad worksheet dimensions.
        header = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        fields = []
        nonempty = 0
        for i, name in enumerate(header):
            if name is None or str(name).strip() == "":
                continue
            nonempty += 1
            hit = classify(name)
            if hit:
                fields.append({"column_1based": i + 1, **hit})
        counts = dict(collections.Counter(x["role"] for x in fields))
        sheets.append({
            "sheet_index_1based": sheet_ix,
            "declared_rows_unverified": declared_rows,
            "declared_columns_unverified": declared_columns,
            "header_row": header_row,
            "nonempty_header_cells": nonempty,
            "matched_known_fields": len(fields),
            "matched_by_role": counts,
            "recognized_fields": fields,
        })
    return sheets


def is_missing(value):
    if value is None:
        return True
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return True
    return isinstance(value, str) and value.strip().upper() in MISSING_TEXT


def subject_key(value):
    if is_missing(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    s = str(value).strip()
    return s if len(s) < 100 else None


def as_number(value):
    if is_missing(value) or isinstance(value, bool):
        return None
    if isinstance(value, (float, int)):
        x = float(value)
    elif isinstance(value, str):
        try:
            x = float(value.strip())
        except ValueError:
            return None
    else:
        return None
    return x if math.isfinite(x) else None


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        val = value.strip()
        if not val:
            return None
        try:
            return dt.date.fromisoformat(val[:10])
        except ValueError:
            pass
        for fmt in ("%m/%d/%Y", "%Y/%m/%d", "%m-%d-%Y"):
            try:
                return dt.datetime.strptime(val, fmt).date()
            except ValueError:
                pass
    return None


def visit_month_and_token(value, date):
    """Return nominal month (if known) and non-identifying visit token.

    In-memory date tokens are never written to output.
    """
    if not is_missing(value):
        v = str(value).strip().lower()
        if v in ("bl", "baseline", "m0", "m00", "m000"):
            return 0, "m0"
        m = re.fullmatch(r"m0*(\d+)", v)
        if m:
            n = int(m.group(1))
            return n, f"m{n}"
        # Screening is distinct from baseline.
        return None, "code:" + v
    if date is not None:
        return None, "date:" + date.isoformat()
    return None, None


def diagnosis_category(raw, field):
    if is_missing(raw):
        return "Unknown"
    s = canonical(raw)
    # DXCHANGE codes are transitions, not current-diagnosis codes.
    if field == "DXCHANGE":
        mapping = {"1": "CN", "2": "MCI", "3": "Dementia", "4": "MCI", "5": "Dementia", "6": "Dementia", "7": "CN", "8": "MCI", "9": "CN"}
        return mapping.get(s, "Other_or_unmapped")
    if s in ("1", "CN", "NL", "NORMAL", "CONTROL", "COGNITIVELYNORMAL", "COGNITIVELYUNIMPAIRED", "CU", "HC", "SMC", "SCD"):
        return "CN_or_subjective_complaint"  # SMC/SCD should be separated during formal inclusion.
    if s in ("2", "MCI", "EMCI", "LMCI", "EARLYMCI", "LATEMCI", "MILDCOGNITIVEIMPAIRMENT"):
        return "MCI"
    if s in ("3", "AD", "DEMENTIA", "ALZHEIMERSDISEASE", "ALZHEIMERDISEASE", "PROBABLEAD"):
        return "Dementia_or_AD"
    return "Other_or_unmapped"


def choose_column(fields, choices):
    # Exact canonical matches only; fallback to baseline columns is intentional only for DXBL.
    for target in choices:
        match = next((x for x in fields if x["field"] == target), None)
        if match:
            return match["column_1based"] - 1, match["field"]
    return (None, None)


def safe_count(n):
    # Output suppression avoids publishing very small cell sizes in group distributions.
    return "<5" if 0 < n < 5 else n


def audit_cohort(ws, fields, header_row, max_rows, log_every):
    index = {}
    index["subject"] = choose_column(fields, ("RID", "PTID", "SUBJECTID", "SUBJECT", "PARTICIPANTID", "SUBJID"))
    index["visit"] = choose_column(fields, ("VISCODE2", "VISCODE", "LONIVISCODE", "VISITCODE", "VISIT"))
    index["date"] = choose_column(fields, ("EXAMDATE", "VISITDATE", "ASSESSDATE", "ASSESSMENTDATE"))
    index["dx"] = choose_column(fields, ("DX", "DXCURREN", "DIAGNOSIS", "DXCHANGE", "DIAGNOSISCODE"))
    index["dxbl"] = choose_column(fields, ("DXBL", "DIAGNOSISBL"))
    for score in SCORES:
        alts = {"MMSE": ("MMSE", "MMSCORE"), "ADAS13": ("ADAS13",), "ADAS11": ("ADAS11",), "CDRSB": ("CDRSB",), "MOCA": ("MOCA",), "FAQ": ("FAQ",)}[score]
        index[score] = choose_column(fields, alts)
    if index["subject"][0] is None:
        return {"status": "cannot_identify_subjects", "reason": "Missing recognized RID/PTID-type column", "selected_fields": {k: val[1] for k, val in index.items()}}

    def at(row, k):
        pos = index[k][0]
        return row[pos] if pos is not None and pos < len(row) else None

    counters = collections.Counter()
    subjects = set()
    subject_visits = collections.defaultdict(set)
    subject_dx_base = {}
    # Longitudinal scores: dict[name] -> dict[internal RID] -> set[(month,visit_token)]
    score_subject_visits = {score: collections.defaultdict(set) for score in SCORES}
    score_subjects = {score: set() for score in SCORES}
    score_rows = collections.Counter()
    score_invalid = collections.Counter()
    score_baseline_subjects = {score: set() for score in SCORES}
    start = time.monotonic()
    exhausted = True
    ws.reset_dimensions()
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        counters["rows_scanned"] += 1
        if max_rows and counters["rows_scanned"] > max_rows:
            counters["rows_scanned"] -= 1
            exhausted = False
            break
        if log_every and counters["rows_scanned"] % log_every == 0:
            print(f"Scanned {counters['rows_scanned']} data rows...", file=sys.stderr, flush=True)
        rid = subject_key(at(row, "subject"))
        if rid is None:
            counters["rows_missing_or_invalid_subject_key"] += 1
            continue
        subjects.add(rid)
        visit_value = at(row, "visit")
        exam_date = parse_date(at(row, "date"))
        month, token = visit_month_and_token(visit_value, exam_date)
        if token is not None:
            subject_visits[rid].add(token)
            counters["rows_with_identifiable_visit_or_date"] += 1
        dxbl_raw = at(row, "dxbl")
        if not is_missing(dxbl_raw):
            subject_dx_base[rid] = diagnosis_category(dxbl_raw, index["dxbl"][1])
        if month == 0 and rid not in subject_dx_base:
            dx_raw = at(row, "dx")
            if not is_missing(dx_raw):
                subject_dx_base[rid] = diagnosis_category(dx_raw, index["dx"][1])

        for score in SCORES:
            val = at(row, score)
            if is_missing(val):
                continue
            x = as_number(val)
            limits = SCORE_LIMITS[score]
            if x is None or not (limits[0] <= x <= limits[1]):
                score_invalid[score] += 1
                continue
            score_rows[score] += 1
            score_subjects[score].add(rid)
            if month == 0:
                score_baseline_subjects[score].add(rid)
            if token is not None:
                score_subject_visits[score][rid].add((month, token))

    dxcounts = collections.Counter(subject_dx_base.values())
    cognition = {}
    for score in SCORES:
        by_person = score_subject_visits[score]
        with_2 = 0
        with_bl_12 = 0
        with_bl_24 = 0
        for rid, obs in by_person.items():
            if len({token for month, token in obs}) >= 2:
                with_2 += 1
            months = {month for month, token in obs if month is not None}
            if 0 in months and 12 in months:
                with_bl_12 += 1
            if 0 in months and 24 in months:
                with_bl_24 += 1
        cognition[score] = {
            "selected_column": index[score][1],
            "rows_with_valid_numeric_value": score_rows[score],
            "unique_subjects_with_valid_value": len(score_subjects[score]),
            "unique_subjects_with_baseline_value": len(score_baseline_subjects[score]),
            "unique_subjects_with_2plus_distinct_timepoints": with_2,
            "unique_subjects_with_bl_and_m12_code": with_bl_12,
            "unique_subjects_with_bl_and_m24_code": with_bl_24,
            "nonempty_but_outside_expected_scale_or_nonnumeric_rows": score_invalid[score],
        }
    result = {
        "status": "ok",
        "complete": exhausted,
        "runtime_seconds": round(time.monotonic() - start, 2),
        "selected_fields": {k: val[1] for k, val in index.items()},
        "rows_scanned": counters["rows_scanned"],
        "rows_missing_or_invalid_subject_key": counters["rows_missing_or_invalid_subject_key"],
        "unique_subjects_in_table": len(subjects),
        "subjects_with_2plus_distinct_visits_or_dates": sum(len(v) >= 2 for v in subject_visits.values()),
        "rows_with_identifiable_visit_or_date": counters["rows_with_identifiable_visit_or_date"],
        "baseline_diagnosis_subjects_total": len(subject_dx_base),
        "baseline_diagnosis_subject_counts_coarse_suppressed": {k: safe_count(v) for k, v in sorted(dxcounts.items())},
        "cognition": cognition,
        "other_recognized_inputs": {
            role: [x["field"] for x in fields if x["role"] == role]
            for role in ("amyloid_or_tau_candidate", "imaging_link", "covariate")
        },
        "important_limits": [
            "Counts are based on this single clinical sheet; no MRI matching was performed.",
            "Baseline and m12/m24 visit-code counts are not guaranteed MRI-anchored 12/24-month outcomes.",
            "DIAGNOSIS classifications are provisional; verify coding with the correct ADNI data dictionary.",
            "SMC/SCD are grouped with CN for coarse counts, not for formal cohort definition.",
            "Biomarker column presence does not prove amyloid positivity or assay availability.",
            "Two distinct assessments may include duplicate/derived visits; manual QC remains necessary.",
            "Missing baseline-only columns are not used as follow-up scores.",
            "Formula cells are read from workbook cached values; absent caches appear missing.",
            "No participant-level records, identifiers, dates, or value samples are exported.",
        ]
    }
    return result


def make_report(args):
    file_path = Path(args.xlsx).expanduser()
    if not file_path.is_file():
        raise FileNotFoundError("Input XLSX was not found")
    started = time.monotonic()
    wb = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    try:
        sheets = scan_schema(wb, args.header_row)
        base = {
            "report": "adni_xlsx_aggregate_audit",
            "script_version": VERSION,
            "mode": args.command,
            "input_type": ".xlsx",
            "sheet_count": len(sheets),
            "schema": sheets,
            "privacy": "Aggregates and whitelisted header names only; no original rows or subject IDs exported",
        }
        if args.command == "cohort":
            sheet_arg = args.sheet
            if sheet_arg is None:
                ranked = sorted(sheets, key=lambda s: (
                    any(x["field"] in ("RID", "PTID") for x in s["recognized_fields"]),
                    s["matched_by_role"].get("cognitive", 0),
                    s["matched_known_fields"],
                ), reverse=True)
                chosen = ranked[0]["sheet_index_1based"] if ranked else 1
            elif str(sheet_arg).isdigit():
                chosen = int(sheet_arg)
            else:
                chosen = wb.sheetnames.index(sheet_arg) + 1
            if chosen < 1 or chosen > len(sheets):
                raise ValueError("Sheet index outside workbook")
            fields = sheets[chosen - 1]["recognized_fields"]
            base["cohort_sheet_index_1based"] = chosen
            base["cohort_audit"] = audit_cohort(wb.worksheets[chosen - 1], fields, args.header_row, args.max_rows, args.log_every)
        base["total_runtime_seconds"] = round(time.monotonic() - started, 2)
        return base
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    for command in ("schema", "cohort"):
        cmd = sub.add_parser(command)
        cmd.add_argument("--xlsx", required=True, help="Local path to .xlsx; never uploaded")
        cmd.add_argument("--out", required=True, help="Local destination for aggregate JSON")
        cmd.add_argument("--header-row", type=int, default=1)
        if command == "cohort":
            cmd.add_argument("--sheet", default=None, help="1-based sheet index or sheet name; default best matching sheet")
            cmd.add_argument("--max-rows", type=int, default=0, help="0 = full worksheet; nonzero yields incomplete report")
            cmd.add_argument("--log-every", type=int, default=2000)
    args = parser.parse_args()
    try:
        report = make_report(args)
        destination = Path(args.out)
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Exports are sensitive even when aggregated; avoid world-readable permissions.
        fd = os.open(str(destination), os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        os.fchmod(fd, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2, ensure_ascii=False)
            f.write("\n")
        print(f"Saved aggregate-only audit to {destination}.")
        if args.command == "cohort":
            r = report.get("cohort_audit", {})
            print(f"Rows scanned: {r.get('rows_scanned', 'n/a')}; subjects: {r.get('unique_subjects_in_table', 'n/a')}; complete: {r.get('complete', 'n/a')}")
    except Exception as exc:
        print(f"Audit failed: {type(exc).__name__}. Check input, sheet selection, and available memory.", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
